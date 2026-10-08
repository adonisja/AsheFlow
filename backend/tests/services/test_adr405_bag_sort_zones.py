"""ADR-405 — the sheet's per-bag sort zone, tested against the export's real layout.

ADR-290 was designed from a photograph. This runs the shipped parsers over a
sheet with the operator's actual structure (six trucks, 20 routes, 339 bags), so
the assumptions are checked rather than inferred.

The fixture keeps the real export's LAYOUT — continuation rows, the
`_x000D_` artefact, the WE123 duplicate bag id — but not its values: anchors,
DSP code, bag numbers and shelf letters were transformed (ADR-489 D2), because
Amazon's export is confidential under the DSP Agreement and this repo is public.
Never replace it with an untransformed export.
"""
from pathlib import Path

import openpyxl  # a declared dependency — never importorskip (ADR-311, ADR-489 D3)
import pytest

from app.services.btr_ingestor import (  # noqa: E402
    _with_sort_zones, parse_bag_labels, parse_ov_zones,
)

SHEET = Path(__file__).resolve().parents[1] / "fixtures" / "btr_sheet_verified.xlsx"


@pytest.fixture(scope="module")
def wb():
    assert SHEET.exists(), f"verified BTR sheet missing at {SHEET}"
    return openpyxl.load_workbook(SHEET, data_only=True)


class TestTheSheetMatchesWhatADR290Assumed:
    def test_one_tab_per_truck_named_btr(self, wb):
        assert all(n.startswith("BTR") for n in wb.sheetnames)
        assert len(wb.sheetnames) == 6

    def test_ov_zone_counts_reconcile_with_ov_count(self, wb):
        """ADR-290 verified this on a photo of one route. Here it is checked on
        every route of every truck: a mismatch means the parser split a cell
        wrongly, or the sheet changed shape."""
        import re
        for name in wb.sheetnames:
            ws = wb[name]
            for r in range(5, ws.max_row + 1):
                if not ws.cell(r, 3).value:
                    continue
                declared = ws.cell(r, 6).value
                parsed = sum(z.ov_count for z in parse_ov_zones(ws.cell(r, 7).value))
                assert parsed == declared, (
                    f"{name} route {ws.cell(r, 3).value}: OV zones sum to "
                    f"{parsed}, sheet says {declared}"
                )

    def test_the_ov_cell_carries_a_carriage_return_artifact(self, wb):
        """`_x000D_\\n`, not a bare newline. `parse_ov_zones` strips it; a parser
        that split on '\\n' alone would leave it stuck to a zone label."""
        raw = wb["BTR43"].cell(5, 7).value
        assert "_x000D_" in raw
        assert [z.zone_label for z in parse_ov_zones(raw)] == ["Q-8.9Z", "Q-8.3T"]


    def test_a_zone_label_may_contain_a_letter_mid_section(self):
        """`Q-2.AZ | 3 OV` — an entry from BTR47/WE36.

        The original pattern required digits between the dash and the trailing
        letter, so this entry matched NOTHING and was dropped in silence: the
        route reported 7 OVs where the sheet said 10. Three OV units would never
        have been seeded, and nobody would have known to look for them.
        """
        zones = parse_ov_zones(
            "Q-2.5Y | 1 OV\nQ-2.AZ | 3 OV\nQ-3.3W | 3 OV"
        )
        assert [(z.zone_label, z.ov_count) for z in zones] == [
            ("Q-2.5Y", 1), ("Q-2.AZ", 3), ("Q-3.3W", 3),
        ]


class TestBagSortZones:
    def test_every_bag_in_the_sheet_has_a_zone(self, wb):
        """339 bags, 0 missing. Recorded so a future sheet that drops the column
        fails this test loudly rather than silently importing nulls."""
        total = missing = 0
        for name in wb.sheetnames:
            ws = wb[name]
            for r in range(5, ws.max_row + 1):
                if ws.cell(r, 8).value:
                    total += 1
                    missing += not ws.cell(r, 9).value
        assert total == 339, f"the fixture changed shape: {total} bags"
        assert missing == 0

    def test_a_bags_zone_is_not_its_routes_ov_zone(self, wb):
        """The two columns are different places. WE1's bags are at Q-9.1E while
        its OVs are at Q-8.9Z and Q-8.3T — reading one as the other would send
        the driver to the wrong shelf."""
        ws = wb["BTR43"]
        bag_zone = ws.cell(5, 9).value
        ov_zones = {z.zone_label for z in parse_ov_zones(ws.cell(5, 7).value)}
        assert bag_zone == "Q-9.1E"
        assert bag_zone not in ov_zones

    def test_two_bags_on_one_route_can_sit_in_different_zones(self, wb):
        """Per BAG, not per route — which is why it cannot live on BTRRoute."""
        ws = wb["BTR43"]
        assert ws.cell(9, 9).value == "Q-9.2E"      # Navy 4046, route WE2
        assert ws.cell(10, 9).value == "Q-9.4E"     # Navy 7858, same route


class TestZipping:
    def test_zones_attach_by_position(self):
        bags = _with_sort_zones(
            parse_bag_labels("Yellow 1029\nYellow 6414"), "Q-9.1E\nQ-9.2E",
        )
        assert [(b.bag_id, b.sort_zone) for b in bags] == [
            ("1029", "Q-9.1E"), ("6414", "Q-9.2E"),
        ]

    def test_a_short_zone_column_leaves_the_rest_null(self):
        """Never shift the alignment. A bag pointed at the wrong shelf is worse
        than a bag with no shelf: the driver walks there and finds someone
        else's tote."""
        bags = _with_sort_zones(parse_bag_labels("Yellow 1029\nYellow 6414"), "Q-9.1E")
        assert bags[0].sort_zone == "Q-9.1E"
        assert bags[1].sort_zone is None

    def test_an_absent_zone_column_is_not_an_error(self):
        """A sheet is an external file. A missing column degrades to null."""
        bags = _with_sort_zones(parse_bag_labels("Yellow 1029"), None)
        assert bags[0].sort_zone is None
